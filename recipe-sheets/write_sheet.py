"""
recipes/*.json を Google スプレッドシートに書き込む。
「一覧」シート + 料理ごとに 1 シート、という構成にする。

使い方:
  python write_sheet.py                      # Google スプレッドシートへ書き込み
  python write_sheet.py --xlsx out.xlsx      # Google を使わず Excel ファイルに出力（動作確認用）

環境変数（Google モード）:
  GOOGLE_SERVICE_ACCOUNT_JSON  サービスアカウントの JSON キー（中身をそのまま）  必須
  RECIPE_SPREADSHEET_ID        書き込み先スプレッドシートの ID。未設定なら新規作成
  SHARE_WITH_EMAIL             新規作成時に編集権限を付与するメールアドレス
"""

import argparse
import json
import os
import re
import sys
import time
from pathlib import Path

from common import RECIPES_DIR, load_json

INDEX_SHEET = "一覧"
INDEX_HEADER = [
    "No", "料理名", "人数", "合計時間(分)", "カロリー/1人前(kcal)",
    "タグ", "チャンネル", "動画の長さ", "動画URL", "レシピシート", "備考",
]

# 行スタイルの種類: title / section / header / muted
Styles = dict[int, str]


# ---------------------------------------------------------------- レイアウト

def sheet_title_for(recipe: dict, used: set[str]) -> str:
    """シート名に使えない文字を除き、100 文字以内・重複なしにする。"""
    name = recipe.get("dish_name") or recipe["source"].get("title") or recipe["source"]["video_id"]
    name = re.sub(r"[\[\]:*?/\\]", " ", name).strip() or recipe["source"]["video_id"]
    name = name[:90]
    base, n = name, 2
    while name in used:
        name = f"{base} ({n})"
        n += 1
    used.add(name)
    return name


def build_recipe_rows(recipe: dict) -> tuple[list[list], Styles]:
    src = recipe["source"]
    rows: list[list] = []
    styles: Styles = {}

    def add(row: list, style: str | None = None) -> None:
        if style:
            styles[len(rows)] = style
        rows.append(row)

    add([recipe.get("dish_name", "")], "title")
    add([recipe.get("summary", "")])
    add([])
    add([
        "人数", recipe.get("servings", ""),
        "準備", f"{recipe.get('prep_time_min', 0)}分",
        "調理", f"{recipe.get('cook_time_min', 0)}分",
        "合計", f"{recipe.get('total_time_min', 0)}分",
        "カロリー(1人前)", f"{recipe.get('calories_per_serving', 0)} kcal",
    ], "header")
    add(["カロリー補足", recipe.get("calories_note", "")], "muted")
    add([])

    add(["■ 動画情報"], "section")
    add(["タイトル", src.get("title", "")])
    add(["チャンネル", src.get("channel", "")])
    add(["動画の長さ", src.get("duration", "")])
    add(["URL", src.get("url", "")])
    add(["タグ", ", ".join(recipe.get("tags") or [])])
    add([])

    add(["■ 材料", recipe.get("servings", "")], "section")
    add(["材料", "分量", "メモ"], "header")
    for ing in recipe.get("ingredients") or []:
        add([ing.get("name", ""), ing.get("amount", ""), ing.get("note", "")])
    add([])

    add(["■ 作り方"], "section")
    add(["No", "手順", "ポイント", "動画の時間"], "header")
    for step in recipe.get("steps") or []:
        add([step.get("order", ""), step.get("text", ""), step.get("point", ""), step.get("timestamp", "")])
    add([])

    add(["■ コツ・ポイント"], "section")
    for tip in recipe.get("tips") or []:
        add(["・", tip])
    add([])

    add(["生成日時", recipe.get("generated_at", ""), "モデル", recipe.get("model", "")], "muted")
    add(["字幕", "あり" if src.get("has_transcript") else "なし（概要欄のみから作成）"], "muted")
    return rows, styles


def build_index_rows(entries: list[tuple[dict, str | None]]) -> list[list]:
    """entries: (recipe, sheet_name or None)"""
    rows = [INDEX_HEADER]
    for i, (recipe, sheet_name) in enumerate(entries, start=1):
        src = recipe["source"]
        rows.append([
            i,
            recipe.get("dish_name", "") if recipe.get("is_recipe") else src.get("title", ""),
            recipe.get("servings", "") if recipe.get("is_recipe") else "",
            recipe.get("total_time_min", "") if recipe.get("is_recipe") else "",
            recipe.get("calories_per_serving", "") if recipe.get("is_recipe") else "",
            ", ".join(recipe.get("tags") or []),
            src.get("channel", ""),
            src.get("duration", ""),
            src.get("url", ""),
            sheet_name or "",
            "" if recipe.get("is_recipe") else f"料理動画ではない: {recipe.get('not_recipe_reason', '')}",
        ])
    return rows


def load_recipes() -> list[dict]:
    files = sorted(RECIPES_DIR.glob("*.json")) if RECIPES_DIR.exists() else []
    return [load_json(p) for p in files]


# ---------------------------------------------------------------- Google Sheets

def write_google(recipes: list[dict]) -> str:
    import gspread

    sa_json = os.environ.get("GOOGLE_SERVICE_ACCOUNT_JSON")
    if not sa_json:
        raise SystemExit("GOOGLE_SERVICE_ACCOUNT_JSON が未設定です（--xlsx で Excel 出力もできます）")
    gc = gspread.service_account_from_dict(json.loads(sa_json))

    spreadsheet_id = os.environ.get("RECIPE_SPREADSHEET_ID", "").strip()
    if spreadsheet_id:
        sh = gc.open_by_key(spreadsheet_id)
    else:
        sh = gc.create("料理動画レシピ集")
        share_with = os.environ.get("SHARE_WITH_EMAIL", "").strip()
        if share_with:
            sh.share(share_with, perm_type="user", role="writer", notify=False)
        print(f"新しいスプレッドシートを作成しました: {sh.url}")
        print("次回以降は RECIPE_SPREADSHEET_ID に次の ID を設定してください:", sh.id)

    def get_or_create(title: str, rows: int, cols: int):
        try:
            ws = sh.worksheet(title)
            ws.clear()
            if ws.row_count < rows or ws.col_count < cols:
                ws.resize(rows=max(ws.row_count, rows), cols=max(ws.col_count, cols))
            return ws
        except gspread.WorksheetNotFound:
            return sh.add_worksheet(title=title, rows=rows, cols=cols)

    def apply_styles(ws, styles: Styles, ncols: int) -> None:
        last_col = chr(ord("A") + min(ncols, 26) - 1)
        fmt = {
            "title": {"textFormat": {"bold": True, "fontSize": 16}},
            "section": {"textFormat": {"bold": True, "fontSize": 12},
                        "backgroundColor": {"red": 0.93, "green": 0.93, "blue": 0.93}},
            "header": {"textFormat": {"bold": True},
                       "backgroundColor": {"red": 0.97, "green": 0.97, "blue": 0.97}},
            "muted": {"textFormat": {"foregroundColor": {"red": 0.5, "green": 0.5, "blue": 0.5}}},
        }
        requests = [
            {"range": f"A{i + 1}:{last_col}{i + 1}", "format": fmt[s]}
            for i, s in styles.items()
        ]
        if requests:
            ws.batch_format(requests)

    # 料理ごとのシート
    used: set[str] = set()
    entries: list[tuple[dict, str | None]] = []
    used.add(INDEX_SHEET)
    for recipe in recipes:
        if not recipe.get("is_recipe"):
            entries.append((recipe, None))
            continue
        title = sheet_title_for(recipe, used)
        rows, styles = build_recipe_rows(recipe)
        ncols = max(len(r) for r in rows)
        ws = get_or_create(title, len(rows) + 5, max(ncols, 10))
        ws.update(range_name="A1", values=rows, value_input_option="RAW")
        apply_styles(ws, styles, ncols)
        ws.columns_auto_resize(0, ncols - 1)
        entries.append((recipe, title))
        print(f"  ✓ シート: {title}")
        time.sleep(1)  # Sheets API の書き込み上限（60回/分）対策

    # 一覧シート（先頭に配置）
    index_rows = build_index_rows(entries)
    for row in index_rows[1:]:
        if row[9]:
            row[9] = f'=HYPERLINK("#gid={sh.worksheet(row[9]).id}", "{row[9]}")'
    try:
        index_ws = sh.worksheet(INDEX_SHEET)
        index_ws.clear()
    except gspread.WorksheetNotFound:
        first = sh.sheet1
        if first.title in ("Sheet1", "シート1") and not spreadsheet_id:
            first.update_title(INDEX_SHEET)
            index_ws = first
        else:
            index_ws = sh.add_worksheet(title=INDEX_SHEET, rows=len(index_rows) + 20, cols=len(INDEX_HEADER))
    index_ws.update(range_name="A1", values=index_rows, value_input_option="USER_ENTERED")
    index_ws.format("A1:K1", {"textFormat": {"bold": True},
                              "backgroundColor": {"red": 0.93, "green": 0.93, "blue": 0.93}})
    index_ws.freeze(rows=1)
    index_ws.columns_auto_resize(0, len(INDEX_HEADER) - 1)
    if index_ws.index != 0:
        index_ws.update_index(0)
    return sh.url


# ---------------------------------------------------------------- Excel (openpyxl)

def write_xlsx(recipes: list[dict], out: Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    fills = {
        "section": PatternFill("solid", fgColor="EDEDED"),
        "header": PatternFill("solid", fgColor="F7F7F7"),
    }
    fonts = {
        "title": Font(bold=True, size=16),
        "section": Font(bold=True, size=12),
        "header": Font(bold=True),
        "muted": Font(color="808080"),
    }

    wb = Workbook()
    index_ws = wb.active
    index_ws.title = INDEX_SHEET

    used: set[str] = {INDEX_SHEET}
    entries: list[tuple[dict, str | None]] = []
    for recipe in recipes:
        if not recipe.get("is_recipe"):
            entries.append((recipe, None))
            continue
        title = sheet_title_for(recipe, used)[:31]  # Excel のシート名は 31 文字まで
        ws = wb.create_sheet(title=title)
        rows, styles = build_recipe_rows(recipe)
        for i, row in enumerate(rows, start=1):
            for j, value in enumerate(row, start=1):
                cell = ws.cell(row=i, column=j, value=value)
                style = styles.get(i - 1)
                if style:
                    cell.font = fonts[style]
                    if style in fills:
                        cell.fill = fills[style]
                if isinstance(value, str) and value.startswith("http"):
                    cell.hyperlink = value
                    cell.font = Font(color="0563C1", underline="single")
        for col, width in zip("ABCD", (18, 60, 40, 12)):
            ws.column_dimensions[col].width = width
        entries.append((recipe, title))

    for i, row in enumerate(build_index_rows(entries), start=1):
        for j, value in enumerate(row, start=1):
            cell = index_ws.cell(row=i, column=j, value=value)
            if i == 1:
                cell.font = fonts["header"]
                cell.fill = fills["section"]
            elif j == 9 and value:
                cell.hyperlink = value
                cell.font = Font(color="0563C1", underline="single")
            elif j == 10 and value:
                cell.hyperlink = f"#'{value}'!A1"
                cell.font = Font(color="0563C1", underline="single")
    for col, width in zip("ABCDEFGHIJK", (5, 30, 10, 12, 16, 24, 20, 10, 44, 30, 30)):
        index_ws.column_dimensions[col].width = width
    index_ws.freeze_panes = "A2"

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


# ---------------------------------------------------------------- main

def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--xlsx", type=Path, help="Google の代わりに Excel ファイルへ出力する")
    args = parser.parse_args()

    recipes = load_recipes()
    if not recipes:
        print("recipes/ にレシピがありません。先に make_recipes.py を実行してください。")
        return 0
    n_recipes = sum(1 for r in recipes if r.get("is_recipe"))
    print(f"{len(recipes)} 本の動画（うち料理動画 {n_recipes} 本）を書き込みます")

    if args.xlsx:
        out = write_xlsx(recipes, args.xlsx)
        print(f"✓ Excel ファイルを出力しました: {out}")
    else:
        url = write_google(recipes)
        print(f"✓ Google スプレッドシートを更新しました: {url}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
